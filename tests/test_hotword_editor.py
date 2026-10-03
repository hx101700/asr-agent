"""验证可编辑热词表格和上下文错误的定位与修正。"""

import json
from datetime import datetime

from openpyxl import Workbook, load_workbook

from asr_runtime.application.inputs import import_hotwords
from asr_runtime.application.rules import ValidationError, validate_context, validate_hotword_rows
from asr_runtime.models import MAX_HOTWORD_ROWS
from asr_runtime.utils.i18n import language_scope
from tests.support import RuntimeTestCase


class HotwordEditorTests(RuntimeTestCase):
    def workbook(self, rows):
        """保存带固定表头的测试工作簿。"""
        path = self.runtime.path("editor.xlsx")
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "热词"
        sheet.append(["text", "weight"])
        for row in rows:
            sheet.append(row)
        workbook.save(path)
        workbook.close()
        return path

    def test_import_preserves_invalid_rows_and_edits_replace_file_input(self):
        """验证非法权重和冲突行在网页修改后即可提交，原Excel保持原样。"""
        path = self.workbook([("术语", 4), ("另一个词", "错误"), ("术语", 5)])
        original = path.read_bytes()
        report = import_hotwords(path)
        self.assertEqual([(issue["row"], issue["field"]) for issue in report["issues"]],
                         [(3, "weight"), (4, "weight")])
        self.assertEqual(report["rows"][1]["weight"], "错误")
        self.assertEqual(path.read_bytes(), original)

        path.unlink()
        report["rows"][1]["weight"] = "4"
        report["rows"][2]["weight"] = "4"
        result = validate_hotword_rows(report["rows"])
        self.assertEqual(result["vocabulary"], {"术语": 4, "另一个词": 4})
        self.assertIn("第4行与第2行", result["warnings"][0])

    def test_formula_errors_belong_to_cells_and_disappear_after_edit(self):
        """验证公式原文及错误定位可见，改写后使用当前单元格值。"""
        report = import_hotwords(self.workbook([("=1+1", 4), ("另一个词", "=2+2")]))
        self.assertEqual(report["rows"], [
            {"row": 2, "text": "=1+1", "weight": 4, "invalid_fields": ["text"]},
            {"row": 3, "text": "另一个词", "weight": "=2+2", "invalid_fields": ["weight"]},
        ])
        self.assertEqual([(issue["row"], issue["field"]) for issue in report["issues"]],
                         [(2, "text"), (3, "weight")])
        self.assertTrue(all("不接受公式" in issue["message"] for issue in report["issues"]))
        report["rows"][0].update(text="术语", invalid_fields=[])
        report["rows"][1].update(weight="4", invalid_fields=[])
        self.assertEqual(validate_hotword_rows(report["rows"])["count"], 2)

    def test_literal_equals_prefix_is_distinct_from_excel_formula(self):
        """验证相同的等号前缀文本仅在Excel确认为公式时需要改写。"""
        path = self.workbook([("=1+1", 4), ("=1+1", 4)])
        workbook = load_workbook(path)
        workbook.active["A3"].data_type = "s"
        workbook.save(path)
        workbook.close()
        report = import_hotwords(path)
        self.assertEqual([(issue["row"], issue["field"]) for issue in report["issues"]], [(2, "text")])
        self.assertNotIn("invalid_fields", report["rows"][1])
        self.assertEqual(validate_hotword_rows([report["rows"][1]])["vocabulary"], {"=1+1": 4})
        # 网页输入保存为固定文本，编辑动作会移除原Excel的类型标记。
        report["rows"][0]["invalid_fields"] = []
        self.assertEqual(validate_hotword_rows(report["rows"])["vocabulary"], {"=1+1": 4})

    def test_limits_report_all_excess_rows_at_once(self):
        """验证总词数和超级词超限时一次返回全部超限行。"""
        for count, weight, expected in ((53, 50, [(52, "weight"), (53, "weight"), (54, "weight")]),
                                         (2002, 4, [(2002, "text"), (2003, "text")])):
            rows = [{"row": index + 2, "text": f"term{index}", "weight": weight} for index in range(count)]
            with self.subTest(count=count, weight=weight), self.assertRaises(ValidationError) as caught:
                validate_hotword_rows(rows)
            self.assertEqual([(issue["row"], issue["field"]) for issue in caught.exception.details], expected)

    def test_dates_and_excel_errors_are_visible_serializable_and_require_edit(self):
        """验证日期和Excel错误保留可见值及格级标记，编辑后可重新校验。"""
        report = import_hotwords(self.workbook([(datetime(2026, 1, 2), 4), ("另一个词", "#VALUE!")]))
        json.dumps(report, allow_nan=False)
        self.assertIn("2026-01-02", report["rows"][0]["text"])
        self.assertEqual(report["rows"][0]["invalid_fields"], ["text"])
        self.assertEqual(report["rows"][1]["weight"], "#VALUE!")
        self.assertEqual(report["rows"][1]["invalid_fields"], ["weight"])
        with self.assertRaises(ValidationError):
            validate_hotword_rows(report["rows"])
        report["rows"][0].update(text="术语", invalid_fields=[])
        report["rows"][1].update(weight="4", invalid_fields=[])
        self.assertEqual(validate_hotword_rows(report["rows"])["count"], 2)

    def test_text_weights_use_explicit_integer_spelling(self):
        """验证直接填写的整数权重可用且错误原值保持可修改。"""
        for weight in (1, 4.0, "1", "4", "50"):
            with self.subTest(weight=weight):
                self.assertEqual(validate_hotword_rows([{"row": 1, "text": "术语", "weight": weight}])["count"], 1)
        for weight in (True, False, "04", "4.0", " 4", "4 ", "5e1", "４", "", None, 6, 2.5):
            row = {"row": 7, "text": "术语", "weight": weight}
            with self.subTest(weight=weight), self.assertRaises(ValidationError) as caught:
                validate_hotword_rows([row])
            self.assertEqual(row["weight"], weight)
            self.assertEqual(caught.exception.details[0]["row"], 7)
            self.assertEqual(caught.exception.details[0]["field"], "weight")

    def test_empty_table_reports_a_cell_and_accepts_new_manual_row(self):
        """验证空模板保留添加词条的操作路径。"""
        report = import_hotwords(self.workbook([]))
        self.assertEqual(report["issues"][0]["field"], "text")
        for rows in ([], [{"row": 9, "text": "", "weight": ""}]):
            with self.subTest(rows=rows), self.assertRaises(ValidationError) as caught:
                validate_hotword_rows(rows)
            self.assertEqual(caught.exception.details[0]["field"], "text")
        result = validate_hotword_rows([{"row": 1, "text": "Kubernetes", "weight": "4"}])
        self.assertEqual(result["vocabulary"], {"Kubernetes": 4})

    def test_row_protocol_rejects_ambiguous_identity_and_non_scalar_values(self):
        """验证错误行协议给出可操作错误并限制表格处理规模。"""
        valid = {"row": 1, "text": "术语", "weight": 4}
        for rows in (None, {}, [valid, valid], [{**valid, "row": True}],
                     [{**valid, "text": ["术语"]}], [{**valid, "invalid_fields": ["other"]}],
                     [valid] * (MAX_HOTWORD_ROWS + 1)):
            with self.subTest(rows_type=type(rows).__name__), self.assertRaises(ValidationError):
                validate_hotword_rows(rows)

    def test_context_messages_name_empty_length_and_character_problems(self):
        """验证参考文本错误说明具体问题并支持中英文。"""
        for language, expected in (("zh-CN", ("为空", "超出 3", "第 2", "U+0000")),
                                   ("en", ("empty", "by 3", "Character 2", "U+0000"))):
            with language_scope(language):
                for value, fragment in ((" \n", expected[0]), ("文" * 403, expected[1]),
                                        ("a\x00b", expected[2]), ("a\x00b", expected[3])):
                    with self.subTest(language=language, fragment=fragment), self.assertRaises(ValidationError) as caught:
                        validate_context(value)
                    self.assertIn(fragment, str(caught.exception))
                    self.assertEqual(caught.exception.field, "context")
        # 官方允许词表、自然语言和混合内容，本机仅检查可传输的文本规格。
        text = "Kubernetes、Bulge Bracket\n补充参考文字🙂"
        self.assertEqual(validate_context(text), text)
