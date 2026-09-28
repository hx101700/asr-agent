import hashlib
import os
import subprocess
import wave
import zipfile
from dataclasses import replace
from unittest.mock import patch

from openpyxl import Workbook

from asr_agent.environment import Project
from asr_agent.media import AudioInfo
from asr_agent.validation import (
    AUDIO_SUFFIXES,
    MAX_UPLOAD_BYTES,
    ValidationError,
    file_fingerprint,
    load_hotwords,
    resolve_input,
    validate_audio,
    validate_context,
)
from tests.support import ProjectTestCase


class ValidationTests(ProjectTestCase):
    def setUp(self):
        super().setUp()
        self.data = self.project.path("data")
        self.data.mkdir()

    def audio(self, channels=1):
        path = self.data / "合成.wav"
        with wave.open(str(path), "wb") as output:
            output.setnchannels(channels)
            output.setsampwidth(2)
            output.setframerate(8000)
            output.writeframes(b"\x00\x00" * 8000 * channels)
        return path

    def hotwords(self, rows, *, headers=("text", "weight"), sheet_name="热词"):
        path = self.data / "合成热词.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = sheet_name
        sheet.append(headers)
        for row in rows:
            sheet.append(row)
        workbook.save(path)
        workbook.close()
        return path

    def test_resolve_accepts_project_relative_regular_file(self):
        path = self.audio()
        self.assertEqual(resolve_input(self.project, "data/合成.wav", AUDIO_SUFFIXES, "audio"), path)

    def test_resolve_rejects_outside_hidden_protected_and_wrong_kind(self):
        self.audio()
        for directory in (".state", "credentials", "node_modules"):
            target = self.project.path(f"{directory}/test.wav")
            target.parent.mkdir(exist_ok=True)
            target.write_bytes(b"synthetic")
        for value in ("", "../outside.wav", ".state/test.wav", "credentials/test.wav",
                      "node_modules/test.wav", "data", "missing.wav", ".env"):
            with self.subTest(value=value), self.assertRaises(ValidationError) as caught:
                resolve_input(self.project, value, AUDIO_SUFFIXES, "audio")
            self.assertEqual(caught.exception.field, "audio")

    def test_resolve_rejects_junction_or_symlink_to_outside_project(self):
        scoped_root = self.project.path("scoped")
        scoped_root.mkdir()
        outside = self.data
        self.audio()
        link = scoped_root / "linked"
        if os.name == "nt":
            result = subprocess.run(["cmd", "/c", "mklink", "/J", str(link), str(outside)],
                                    capture_output=True, creationflags=subprocess.CREATE_NO_WINDOW)
            self.assertEqual(result.returncode, 0)
        else:
            link.symlink_to(outside, target_is_directory=True)
        with self.assertRaises(ValidationError):
            resolve_input(Project(scoped_root), "linked/合成.wav", AUDIO_SUFFIXES, "audio")

    def test_resolve_rejects_actual_hidden_attribute_on_windows(self):
        if os.name != "nt":
            return
        import ctypes
        path = self.audio()
        self.assertTrue(ctypes.windll.kernel32.SetFileAttributesW(str(path), 2))
        try:
            with self.assertRaises(ValidationError):
                resolve_input(self.project, path, AUDIO_SUFFIXES, "audio")
        finally:
            ctypes.windll.kernel32.SetFileAttributesW(str(path), 128)

    def test_fingerprint_detects_same_size_content_change(self):
        path = self.data / "test.txt"
        path.write_bytes(b"abc")
        first = file_fingerprint(path)
        self.assertEqual(first["sha256"], hashlib.sha256(b"abc").hexdigest())
        path.write_bytes(b"def")
        second = file_fingerprint(path)
        self.assertEqual(first["size_bytes"], second["size_bytes"])
        self.assertNotEqual(first["sha256"], second["sha256"])

    def test_audio_mono_uses_original_and_stereo_only_plans_conversion(self):
        path = self.audio()
        mono = validate_audio(self.project, path, True)
        self.assertFalse(mono["requires_mono"])
        self.assertEqual(mono["metadata"]["duration_seconds"], 1.0)
        path = self.audio(channels=2)
        before = path.read_bytes()
        stereo = validate_audio(self.project, path, True)
        self.assertTrue(stereo["requires_mono"])
        self.assertIn("保留原文件", stereo["warnings"][0])
        self.assertFalse(validate_audio(self.project, path, False)["requires_mono"])
        self.assertEqual(path.read_bytes(), before)
        self.assertEqual(list(self.data.iterdir()), [path])

    def test_audio_rejects_empty_corrupt_and_unsupported_extension(self):
        for name, contents in (("empty.wav", b""), ("broken.wav", b"not a wave"),
                               ("file.txt", b"not audio")):
            path = self.data / name
            path.write_bytes(contents)
            with self.subTest(name=name), self.assertRaises(ValidationError):
                validate_audio(self.project, path, True)

    def test_audio_rejects_same_length_modification_during_probe(self):
        path = self.audio()
        before = path.stat()

        def changing_probe(source):
            content = source.read_bytes()
            source.write_bytes(content[:-2] + b"\x01\x01")
            os.utime(source, ns=(before.st_atime_ns, before.st_mtime_ns + 1_000_000))
            return AudioInfo(1, 8000, 1.0, before.st_size, "wav", 1)

        with patch("asr_agent.validation.probe_audio", side_effect=changing_probe):
            with self.assertRaisesRegex(ValidationError, "校验期间发生变化"):
                validate_audio(self.project, path, True)

    def test_audio_duration_and_actual_container_boundaries(self):
        path = self.audio()
        normal = AudioInfo(1, 8000, 1.0, path.stat().st_size, "wav", 1)
        for duration in (None, 0, -1, float("nan"), float("inf"), 43200.001):
            with self.subTest(duration=duration), patch("asr_agent.validation.probe_audio",
                    return_value=replace(normal, duration_seconds=duration)):
                with self.assertRaises(ValidationError):
                    validate_audio(self.project, path, True)
        with patch("asr_agent.validation.probe_audio", return_value=replace(normal, format_name="aiff")):
            with self.assertRaisesRegex(ValidationError, "实际媒体格式"):
                validate_audio(self.project, path, False)
        with patch("asr_agent.validation.probe_audio",
                   return_value=replace(normal, duration_seconds=43200, audio_tracks=2)):
            result = validate_audio(self.project, path, True)
            self.assertEqual(len(result["warnings"]), 2)
            self.assertTrue(any("索引0" in text for text in result["warnings"]))

    def test_upload_size_applies_to_actual_upload_not_source_to_be_merged(self):
        path = self.audio(channels=2)
        big = AudioInfo(2, 8000, 1.0, MAX_UPLOAD_BYTES + 1, "wav", 1)
        with patch("asr_agent.validation.probe_audio", return_value=big):
            with self.assertRaisesRegex(ValidationError, "1 GB"):
                validate_audio(self.project, path, False)
            with patch("asr_agent.validation.file_fingerprint",
                       return_value={"size_bytes": big.size_bytes, "mtime_ns": path.stat().st_mtime_ns,
                                     "sha256": "synthetic"}):
                self.assertTrue(validate_audio(self.project, path, True)["requires_mono"])

    def test_context_counts_unicode_characters_and_preserves_text(self):
        text = "🙂" * 398 + " \n"
        self.assertEqual(validate_context(text), text)
        for invalid in ("", " \n", "测" * 401, "a\x00b", "a\ud800", None):
            with self.subTest(length=len(invalid) if invalid is not None else None):
                with self.assertRaises(ValidationError) as caught:
                    validate_context(invalid)
                self.assertEqual(caught.exception.field, "context")

    def test_hotwords_accepts_chinese_headers_and_reports_blank_and_duplicate_rows(self):
        path = self.hotwords([("语音实验室", 4), (None, None), ("语音实验室", 4), ("hello world", 2)],
                             headers=("热词", "权重"))
        before = path.read_bytes()
        result = load_hotwords(self.project, path)
        self.assertEqual(result["vocabulary"], {"语音实验室": 4, "hello world": 2})
        self.assertEqual(result["ignored_blank_rows"], 1)
        self.assertEqual(len(result["warnings"]), 3)
        self.assertEqual(path.read_bytes(), before)

    def test_hotwords_reports_formula_conflict_empty_word_and_invalid_weights_by_row(self):
        path = self.hotwords([
            ("first", 4), ("first", 3), ("=1+1", 4), (None, 2), ("word", True),
            ("second", 2.5), ("third", "4"), (" 热词", 3), ("tab\tword", 2),
        ])
        with self.assertRaises(ValidationError) as caught:
            load_hotwords(self.project, path)
        details = caught.exception.details
        self.assertEqual({error["row"] for error in details}, set(range(3, 11)))
        self.assertTrue(all(set(error) == {"row", "field", "message"} for error in details))
        self.assertEqual(caught.exception.field, "hotwords_path")

    def test_hotword_length_rules(self):
        valid = self.hotwords([("汉" * 15, 1), ("a b c d e f g", 5)])
        self.assertEqual(load_hotwords(self.project, valid)["count"], 2)
        invalid = self.hotwords([("汉" * 16, 1), ("a b c d e f g h", 5)])
        with self.assertRaises(ValidationError) as caught:
            load_hotwords(self.project, invalid)
        self.assertEqual([error["row"] for error in caught.exception.details], [2, 3])

    def test_super_words_are_explicit_and_limited(self):
        path = self.hotwords([(f"term{i}", 50) for i in range(50)])
        with self.assertRaises(ValidationError):
            load_hotwords(self.project, path)
        self.assertEqual(load_hotwords(self.project, path, allow_super=True)["super_count"], 50)
        path = self.hotwords([(f"term{i}", 50) for i in range(51)])
        with self.assertRaises(ValidationError) as caught:
            load_hotwords(self.project, path, allow_super=True)
        self.assertEqual(caught.exception.details[0]["row"], 52)

    def test_hotword_count_limit(self):
        path = self.hotwords([(f"term{i}", 4) for i in range(2000)])
        self.assertEqual(load_hotwords(self.project, path)["count"], 2000)
        path = self.hotwords([(f"term{i}", 4) for i in range(2001)])
        with self.assertRaises(ValidationError) as caught:
            load_hotwords(self.project, path)
        self.assertEqual(caught.exception.details[0]["row"], 2002)

    def test_hotwords_rejects_invalid_header_empty_extra_columns_and_corrupt_file(self):
        for headers, rows in [(("word", "weight"), [("hello", 4)]),
                              (("text", "weight"), []),
                              (("text", "weight", "extra"), [("hello", 4, "x")])]:
            with self.subTest(headers=headers, row_count=len(rows)), self.assertRaises(ValidationError):
                load_hotwords(self.project, self.hotwords(rows, headers=headers))
        path = self.data / "bad.xlsx"
        path.write_text("not a spreadsheet")
        with self.assertRaises(ValidationError):
            load_hotwords(self.project, path)

    def test_hotwords_bounds_archive_size_and_decompressed_size(self):
        path = self.hotwords([("test", 4)])
        with patch("asr_agent.validation.MAX_XLSX_BYTES", 1):
            with self.assertRaisesRegex(ValidationError, "文件上限"):
                load_hotwords(self.project, path)
        with patch("asr_agent.validation.MAX_XLSX_UNCOMPRESSED_BYTES", 1):
            with self.assertRaisesRegex(ValidationError, "解压内容"):
                load_hotwords(self.project, path)

    def test_hotwords_rejects_xml_entities(self):
        path = self.hotwords([("test", 4)])
        with zipfile.ZipFile(path) as original:
            entries = {name: original.read(name) for name in original.namelist()}
        entries["xl/worksheets/sheet1.xml"] = (
            b'<!DOCTYPE worksheet [<!ENTITY word "expanded">]>'
            b'<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
            b'<sheetData><row r="1"><c r="A1" t="inlineStr"><is><t>&word;</t></is></c>'
            b'</row></sheetData></worksheet>'
        )
        with zipfile.ZipFile(path, "w") as modified:
            for name, contents in entries.items():
                modified.writestr(name, contents)
        with self.assertRaises(ValidationError):
            load_hotwords(self.project, path)

    def test_hotwords_rejects_ambiguous_sheets_and_discloses_explicit_selection(self):
        path = self.data / "sheets.xlsx"
        workbook = Workbook()
        workbook.active.title = "one"
        workbook.create_sheet("two")
        workbook.save(path)
        with self.assertRaisesRegex(ValidationError, "多个工作表"):
            load_hotwords(self.project, path)
        sheet = workbook.create_sheet("热词")
        sheet.append(["text", "weight"])
        sheet.append(["test", 4])
        workbook.save(path)
        workbook.close()
        result = load_hotwords(self.project, path)
        self.assertEqual(result["sheet"], "热词")
        self.assertTrue(any("其他工作表" in warning for warning in result["warnings"]))
