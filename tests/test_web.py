import http.client
import io
import json
import threading
import wave
from pathlib import Path
from concurrent.futures import ThreadPoolExecutor
from urllib.parse import quote
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from asr_transcription.application.rules import ValidationError
from asr_transcription.web import create_server
from asr_transcription.application.session import Session, output_directory
from tests.support import RuntimeTestCase


class WebFixture(RuntimeTestCase):
    def setUp(self):
        """准备可供会话上传的合成音频。"""
        super().setUp()
        self.audio = self.runtime.path("data/audio/录音.wav")
        self.audio.parent.mkdir(parents=True)
        with wave.open(str(self.audio), "wb") as audio:
            audio.setnchannels(2)
            audio.setsampwidth(2)
            audio.setframerate(16000)
            audio.writeframes(b"\0" * 64000)
    def payload_for(self, session):
        """上传合成音频并构造默认网页配置。"""
        content = self.audio.read_bytes()
        uploaded = session.upload("audio", self.audio.name, io.BytesIO(content), len(content))
        return {"auth_mode": "console", "audio_upload_id": uploaded["upload_id"],
                        "diarization_enabled": True, "enhancement_mode": "none",
                        "hotwords_upload_id": "", "context": "", "json_directory": "default",
                        "document_directory": "default"}


class SessionTests(WebFixture):
    def setUp(self):
        """准备会话及已上传音频的默认配置。"""
        super().setUp()
        self.session = Session(self.runtime)
        self.payload = self.payload_for(self.session)

    def test_preview_performs_no_conversion_or_job_write(self):
        """验证预览生成内存快照和用户核对信息。"""
        result = self.session.validate(self.payload)
        self.assertEqual(result["summary"]["audio"]["channels"], 2)
        self.assertTrue(any("单声道" in warning for warning in result["summary"]["warnings"]))
        self.assertFalse(self.runtime.path(".state/jobs").exists())
        self.assertFalse(self.runtime.output_root.exists())

    def test_confirm_is_idempotent_even_with_concurrent_requests(self):
        """验证并发确认共用同一配置回执。"""
        preview = self.session.validate(self.payload)
        with ThreadPoolExecutor(max_workers=2) as pool:
            receipts = list(pool.map(self.session.confirm, [preview["validation_id"]] * 2))
        self.assertEqual(receipts[0], receipts[1])
        self.assertEqual(len(list(self.runtime.path(".state/jobs").iterdir())), 1)
        with open(receipts[0]["config_path"], encoding="utf-8") as file:
            config = json.load(file)
        self.assertFalse(config["execution_authorized"])
        self.assertFalse(receipts[0]["execution_started"])
        self.assertEqual(config["region"], "cn-beijing")

    def test_failed_revalidation_invalidates_old_preview(self):
        """验证再次校验失败使旧预览失效。"""
        previous = self.session.validate(self.payload)
        with self.assertRaises(ValidationError):
            self.session.validate({**self.payload, "audio_upload_id": "missing"})
        with self.assertRaises(ValidationError):
            self.session.confirm(previous["validation_id"])

    def test_file_changed_after_preview_cannot_be_confirmed(self):
        """验证预览后文件变化时拒绝确认。"""
        preview = self.session.validate(self.payload)
        uploaded = Path(self.session.uploaded(self.payload["audio_upload_id"], "audio")["path"])
        uploaded.write_bytes(uploaded.read_bytes() + b"changed")
        with self.assertRaises(ValidationError):
            self.session.confirm(preview["validation_id"])
        self.assertFalse(self.runtime.path(".state/jobs").exists())

    def test_confirmation_uses_metadata_without_reading_audio_again(self):
        """验证确认通过文件元信息复用预览快照。"""
        preview = self.session.validate(self.payload)
        with patch("asr_transcription.application.inputs.file_fingerprint", side_effect=AssertionError("must not hash again")), \
             patch("asr_transcription.application.inputs.probe_audio", side_effect=AssertionError("must not probe again")):
            self.assertTrue(self.session.confirm(preview["validation_id"])["ok"])

    def test_api_key_is_not_written_to_config(self):
        """验证确认配置与凭据保持分离。"""
        secret = "synthetic-local-key"
        self.runtime.path(".env").write_text("DASHSCOPE_API_KEY=" + secret, encoding="utf-8")
        preview = self.session.validate({**self.payload, "auth_mode": "api_key"})
        receipt = self.session.confirm(preview["validation_id"])
        with open(receipt["config_path"], encoding="utf-8") as file:
            text = file.read()
        self.assertNotIn(secret, text)
        self.assertNotIn(secret, json.dumps(preview))

    def test_api_key_change_does_not_invalidate_configuration(self):
        """验证Key改变后配置仍能正常确认。"""
        path = self.runtime.path(".env")
        path.write_text("DASHSCOPE_API_KEY=synthetic-key", encoding="utf-8")
        preview = self.session.validate({**self.payload, "auth_mode": "api_key"})
        path.write_text("DASHSCOPE_API_KEY=another-synthetic-key", encoding="utf-8")
        with patch("asr_transcription.application.session.read_api_key", side_effect=AssertionError("must not read credentials")):
            self.assertTrue(self.session.confirm(preview["validation_id"])["ok"])

    def test_description_and_configuration_do_not_read_dotenv(self):
        """验证页面说明与配置流程独立于运行凭据。"""
        with patch("asr_transcription.application.session.read_api_key", side_effect=AssertionError("must not read credentials")):
            self.assertNotIn("auth", self.session.description())
            preview = self.session.validate({**self.payload, "auth_mode": "api_key"})
            self.assertTrue(self.session.confirm(preview["validation_id"])["ok"])

    def test_confirmation_uses_hotwords_snapshot_after_excel_is_removed(self):
        """验证Excel删除后确认仍使用已解析热词快照。"""
        path = self.runtime.path("data/热词.xlsx")
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "热词"
        sheet.append(["text", "weight"])
        sheet.append(["术语", 50])
        workbook.save(path)
        workbook.close()
        content = path.read_bytes()
        uploaded = self.session.upload("hotwords", path.name, io.BytesIO(content), len(content))
        payload = {**self.payload, "enhancement_mode": "both", "hotwords_upload_id": uploaded["upload_id"],
                   "context": "会议涉及测试术语"}
        preview = self.session.validate(payload)
        self.assertEqual(preview["summary"]["enhancement"]["count"], 1)
        self.assertEqual(preview["summary"]["enhancement"]["context_chars"], len(payload["context"]))
        copy = Path(self.session.uploaded(uploaded["upload_id"], "hotwords")["path"])
        copy.unlink()
        with patch("asr_transcription.application.session.load_hotwords", side_effect=AssertionError("must not parse Excel again")):
            receipt = self.session.confirm(preview["validation_id"])
        config = json.loads(Path(receipt["config_path"]).read_text(encoding="utf-8"))
        self.assertEqual(config["enhancement"]["hotwords"]["vocabulary"], {"术语": 50})

    def test_browser_rejects_unapproved_output_paths(self):
        """验证客户端提交源码或其他路径时被拒绝。"""
        for path in ("data/audio", "../outside", "transcriptions/meeting", "transcriptions/../.git", "transcriptions/.private"):
            with self.subTest(path=path), self.assertRaises(ValidationError):
                output_directory(self.runtime, path, "json_directory")

    def test_unknown_fields_and_multiple_audio_values_are_rejected(self):
        """验证未知字段和多音频输入被拒绝。"""
        with self.assertRaises(ValidationError):
            self.session.validate({**self.payload, "model": "other-model"})
        with self.assertRaises(ValidationError):
            self.session.validate({**self.payload, "audio_upload_id": ["a", "b"]})

    def test_browser_source_is_not_read_again_after_upload(self):
        """验证上传后预览使用本机会话副本。"""
        preview = self.session.validate(self.payload)
        self.audio.write_bytes(b"source changed outside local upload")
        receipt = self.session.confirm(preview["validation_id"])
        self.assertTrue(receipt["ok"])

    def test_reselection_invalidates_preview_and_replaces_only_own_copy(self):
        """验证重选使预览失效并替换当前会话副本。"""
        preview = self.session.validate(self.payload)
        old_copy = Path(self.session.uploaded(self.payload["audio_upload_id"], "audio")["path"])
        content = self.audio.read_bytes()
        self.session.upload("audio", "another.wav", io.BytesIO(content), len(content))
        self.assertFalse(old_copy.exists())
        self.assertTrue(self.audio.is_file())
        with self.assertRaises(ValidationError):
            self.session.confirm(preview["validation_id"])

    def test_incomplete_upload_removes_partial_copy(self):
        """验证上传长度异常时清理临时副本。"""
        with self.assertRaises(ValidationError):
            self.session.upload("audio", "short.wav", io.BytesIO(b"short"), 10)
        self.assertFalse(list(self.session.upload_directory.glob("*.part")))

    def test_upload_id_cannot_be_used_by_another_session(self):
        """验证跨会话上传编号被拒绝。"""
        with self.assertRaises(ValidationError):
            Session(self.runtime).validate(self.payload)


class WebServerTests(WebFixture):
    def setUp(self):
        """启动本机HTTP测试服务并准备会话令牌。"""
        super().setUp()
        self.server = create_server(self.runtime)
        self.payload = self.payload_for(self.server.session)
        self.thread = threading.Thread(target=self.server.serve_forever, daemon=True)
        self.thread.start()
        self.origin = f"http://127.0.0.1:{self.server.server_port}"

    def tearDown(self):
        """关闭本机测试服务并清理临时项目。"""
        self.server.shutdown()
        self.server.server_close()
        self.thread.join(timeout=5)
        super().tearDown()

    def request(self, method, path, payload=None, *, token=True, headers=None, body=None):
        """向本机服务发送指定请求并解析响应。"""
        connection = http.client.HTTPConnection("127.0.0.1", self.server.server_port, timeout=5)
        supplied = {"Origin": self.origin}
        if token:
            supplied["X-ASR-Token"] = self.server.session.token
        if headers:
            supplied.update(headers)
        if payload is not None:
            body = json.dumps(payload, ensure_ascii=False).encode("utf-8")
            supplied["Content-Type"] = "application/json"
        connection.request(method, path, body=body, headers=supplied)
        response = connection.getresponse()
        result = response.status, dict(response.getheaders()), response.read()
        connection.close()
        return result

    def test_static_page_has_readonly_key_display_and_does_not_expose_files(self):
        """验证页面提供只读Key控件并限制文件访问入口。"""
        status, headers, body = self.request("GET", "/", token=False)
        self.assertEqual(status, 200)
        self.assertIn(b'id="api-key-value" type="password" readonly', body)
        self.assertNotIn(b'name="api_key"', body)
        self.assertIn("frame-ancestors 'none'", headers["Content-Security-Policy"])
        self.assertEqual(self.request("GET", "/.env")[0], 404)

    def test_token_host_and_origin_are_enforced(self):
        """验证令牌、主机与来源限制。"""
        self.assertEqual(self.request("GET", "/api/session", token=False)[0], 403)
        self.assertEqual(self.request("GET", "/api/session", headers={"X-ASR-Token": "é"})[0], 403)
        self.assertEqual(self.request("GET", "/api/session", headers={"Host": "evil.example"})[0], 403)
        self.assertEqual(self.request("POST", "/api/validate",
                                      headers={"Origin": "https://evil.example"})[0], 403)

    def test_second_server_cannot_reuse_live_port(self):
        """验证已占用端口拒绝第二个服务绑定。"""
        with self.assertRaises(OSError):
            create_server(self.runtime, self.server.server_port)

    def test_cookie_restores_receipt_without_browser_storage(self):
        """验证Cookie会话恢复已保存回执。"""
        status, headers, _ = self.request("GET", "/api/session")
        self.assertEqual(status, 200)
        cookie = headers["Set-Cookie"]
        self.assertIn("HttpOnly", cookie)
        self.assertIn("SameSite=Strict", cookie)
        self.assertNotIn("Max-Age", cookie)
        preview = json.loads(self.request("POST", "/api/validate", self.payload)[2])
        receipt = json.loads(self.request("POST", "/api/confirm", {"validation_id": preview["validation_id"]})[2])
        status, _, body = self.request("GET", "/api/session", token=False, headers={"Cookie": cookie.split(";", 1)[0]})
        self.assertEqual(status, 200)
        self.assertEqual(json.loads(body)["confirmed"]["job_id"], receipt["job_id"])

    def test_confirm_emits_job_receipt_for_codex_without_private_inputs(self):
        """验证确认输出可供Codex读取的脱敏回执。"""
        secret = "synthetic-secret-key"
        context = "只用于测试的内部会议术语"
        self.runtime.path(".env").write_text("DASHSCOPE_API_KEY=" + secret, encoding="utf-8")
        payload = {**self.payload, "auth_mode": "api_key", "enhancement_mode": "context", "context": context}
        preview = json.loads(self.request("POST", "/api/validate", payload)[2])

        with patch("builtins.print") as printed:
            status, _, body = self.request("POST", "/api/confirm", {"validation_id": preview["validation_id"]})

        self.assertEqual(status, 200)
        receipt = json.loads(body)
        event = json.loads(printed.call_args.args[0])
        self.assertEqual(event, {"event": "configured", **receipt})
        self.assertEqual(event["auth_mode"], "api_key")
        self.assertFalse(event["execution_started"])
        self.assertNotIn(secret, printed.call_args.args[0])
        self.assertNotIn(context, printed.call_args.args[0])

    def test_direct_upload_and_template(self):
        """验证直接上传与热词模板下载。"""
        status, _, body = self.request("POST", "/api/upload-audio", body=self.audio.read_bytes(),
                                      headers={"Content-Type": "application/octet-stream",
                                               "X-File-Name": quote(self.audio.name)})
        self.assertEqual(status, 200)
        uploaded = json.loads(body)
        self.assertEqual(uploaded["name"], self.audio.name)
        self.assertEqual(self.request("POST", "/api/validate",
                                      {**self.payload, "audio_upload_id": uploaded["upload_id"]})[0], 200)
        self.assertEqual(self.request("GET", "/api/files?directory=data/audio")[0], 404)
        status, _, body = self.request("GET", "/api/hotwords-template")
        self.assertEqual(status, 200)
        workbook = load_workbook(io.BytesIO(body))
        self.assertEqual(workbook["热词"]["A1"].value, "text")
        workbook.close()

    def test_upload_filename_cannot_supply_a_path(self):
        """验证路径形式的上传文件名被拒绝。"""
        status, _, body = self.request("POST", "/api/upload-audio",
                                   headers={"Content-Type": "application/octet-stream",
                                            "X-File-Name": quote("../escape.wav")})
        self.assertEqual(status, 422)
        self.assertIn("所选文件名或格式不符合要求", json.loads(body)["error"])
        self.assertFalse(self.runtime.path("escape.wav").exists())

    def test_api_key_error_does_not_return_contents(self):
        """验证Key读取失败返回脱敏错误。"""
        status, _, body = self.request("POST", "/api/api-key", {})
        self.assertEqual(status, 422)
        self.assertEqual(json.loads(body)["field"], "auth_mode")
        self.assertEqual(self.request("POST", "/api/auth-status", {"auth_mode": "api_key"})[0], 404)
