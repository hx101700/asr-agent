"""验证三种成品的文字保真、格式和回读检查。"""

import html
import re
import string
from datetime import timedelta
from unittest.mock import patch
from zipfile import ZipFile

from docx import Document
from docx.document import Document as WordDocument
from openpyxl import load_workbook
from openpyxl.workbook import Workbook

from asr_runtime.utils.documents import DocumentError, write_docx, write_markdown, write_xlsx
from asr_runtime.models import Sentence, Transcript
from tests.support import RuntimeTestCase


def sample(*texts):
    """构造包含中文、空句和合成标签的转写样本。"""
    values = texts or ("中文测试。第二句带数字123。", "=SUM(1,2)", "")
    return Transcript(tuple(
        Sentence(index, 1 if index == 1 else 2, 0 if index == 1 else None,
                 (index - 1) * 3_600_000 + 123, index * 3_600_000 + 789,
                 0 if index == 1 else None, text)
        for index, text in enumerate(values, 1)
    ), audio_tracks=2, json_bytes=100, sha256="a" * 64)


class DocumentTests(RuntimeTestCase):
    def output(self, extension):
        """取得指定格式的测试输出路径。"""
        return self.runtime.path(f"transcription.{extension}")

    def write(self, writer, extension, transcript=None):
        """用指定写入器导出样本并返回文件位置。"""
        target = self.output(extension)
        writer(transcript or sample(), target, source_name="合成测试录音.wav", job_id="a" * 32)
        return target

    def test_xlsx_roundtrip_preserves_all_columns_and_has_no_formulas(self):
        """验证Excel回读保留所有列并将正文存为文本。"""
        path = self.write(write_xlsx, "xlsx")
        with path.open("rb") as stream:
            workbook = load_workbook(stream)
            try:
                sheet = workbook["转写明细"]
                self.assertEqual(sheet.max_row, 6)
                self.assertEqual(sheet.freeze_panes, "B4")
                self.assertEqual(sheet.auto_filter.ref, "A3:F6")
                self.assertEqual(sheet["A4"].value, 1)
                self.assertEqual(sheet["B4"].value, "音轨0")
                self.assertEqual(sheet["B5"].value, "未标注（第2组）")
                self.assertEqual(sheet["C4"].value, timedelta(milliseconds=123))
                self.assertEqual(sheet["D4"].value, timedelta(milliseconds=3_600_789))
                self.assertEqual(sheet["E4"].value, "说话人0")
                self.assertEqual(sheet["E5"].value, "未标注")
                self.assertEqual(sheet["F5"].value, "=SUM(1,2)")
                self.assertEqual(sheet["F5"].data_type, "s")
                self.assertIsNone(sheet["F5"].hyperlink)
                self.assertIsNone(sheet["F6"].value)
                self.assertTrue(sheet["F4"].alignment.wrap_text)
                self.assertEqual(sheet.row_dimensions[4].height, 30)
                self.assertEqual(sheet["A3"].value, "序号")
            finally:
                workbook.close()

    def test_xlsx_rejects_text_beyond_cell_storage_limit(self):
        """验证Excel拒绝超过单元格存储上限的正文。"""
        with self.assertRaisesRegex(DocumentError, "32767"):
            self.write(write_xlsx, "xlsx", sample("字" * 32768))

    def test_xlsx_keeps_long_sentence_with_wrapping(self):
        """验证Excel长句保留全文并自动换行。"""
        text = '=HYPERLINK("https://example.invalid") ' + "长句" * 800
        path = self.write(write_xlsx, "xlsx", sample(text))
        with path.open("rb") as stream:
            workbook = load_workbook(stream)
            try:
                sheet = workbook.active
                self.assertEqual(sheet["F4"].value, text)
                self.assertEqual(sheet["F4"].data_type, "s")
                self.assertIsNone(sheet["F4"].hyperlink)
                self.assertEqual(sheet.row_dimensions[4].height, 409)
                self.assertTrue(sheet["F4"].alignment.wrap_text)
                self.assertEqual(sheet["A3"].value, "序号")
            finally:
                workbook.close()
        with ZipFile(path) as archive:
            self.assertFalse(any(name.startswith("xl/externalLinks/") for name in archive.namelist()))

    def test_xlsx_rejects_xml_control_character_without_echoing_content(self):
        """验证Excel控制字符返回脱敏格式错误。"""
        secret_text = "合成敏感测试\x00内容"
        with self.assertRaises(DocumentError) as error:
            self.write(write_xlsx, "xlsx", sample(secret_text))
        self.assertNotIn(secret_text, str(error.exception))

    def test_xlsx_preserves_newlines_carriage_returns_and_unicode(self):
        """验证Excel保留换行、回车和Unicode字符。"""
        text = "  中文 & < > 😀\t测试\r\n下一行\r结尾 "
        path = self.write(write_xlsx, "xlsx", sample(text))
        with path.open("rb") as stream:
            workbook = load_workbook(stream)
            try:
                self.assertEqual(workbook.active["F4"].value, text)
            finally:
                workbook.close()

    def test_xlsx_readback_rejects_text_corruption(self):
        """验证Excel回读能发现正文被改写。"""
        save = Workbook.save

        def corrupt(workbook, path):
            """改写Excel正文以模拟写入后损坏。"""
            save(workbook, path)
            edited = load_workbook(path)
            edited.active["F4"] = "合成篡改"
            save(edited, path)
            edited.close()

        with patch.object(Workbook, "save", corrupt), self.assertRaisesRegex(DocumentError, "回读"):
            self.write(write_xlsx, "xlsx")

    def test_word_roundtrip_retains_exact_text_tags_and_styles(self):
        """验证Word回读保留正文、标签和字体格式。"""
        text = "  原文\t测试\r\n下一行\r单独回车 & <标签> 😀 "
        path = self.write(write_docx, "docx", sample(text, ""))
        document = Document(path)
        self.assertEqual(len(document.paragraphs), 7)
        self.assertEqual(document.paragraphs[0].text, "合成测试录音 录音转写")
        self.assertEqual(document.paragraphs[0].style.name, "Title")
        self.assertEqual(document.paragraphs[4].text, text)
        self.assertEqual(document.paragraphs[6].text, "")
        self.assertIn("00:00:00.123 — 01:00:00.789", document.paragraphs[3].text)
        self.assertIn("说话人0  音轨0", document.paragraphs[3].text)
        self.assertIn("未标注  未标注（第2组）", document.paragraphs[5].text)
        self.assertTrue(document.styles["Transcript Marker"].paragraph_format.keep_with_next)
        self.assertFalse(document.styles["Normal"].paragraph_format.keep_together)
        self.assertEqual(document.sections[0].page_width.inches, 8.5)
        self.assertEqual(document.sections[0].page_height.inches, 11)
        with ZipFile(path) as archive:
            xml = archive.read("word/document.xml")
            self.assertNotIn(b"<w:hyperlink", xml)
            self.assertNotIn(b"<w:pBdr", xml)

    def test_word_retains_long_text_without_excel_limits(self):
        """验证Word完整保存超长正文。"""
        text = "长文本验证" * 8000
        path = self.write(write_docx, "docx", sample(text))
        self.assertEqual(Document(path).paragraphs[-1].text, text)

    def test_word_rejects_xml_control_character_without_echoing_content(self):
        """验证Word控制字符返回脱敏格式错误。"""
        text = "合成测试\x07不可显示字符"
        with self.assertRaises(DocumentError) as error:
            self.write(write_docx, "docx", sample(text))
        self.assertNotIn(text, str(error.exception))

    def test_word_readback_rejects_text_corruption(self):
        """验证Word回读能发现正文被改写。"""
        save = WordDocument.save

        def corrupt(document, path):
            """改写Word正文以模拟写入后损坏。"""
            save(document, path)
            edited = Document(path)
            edited.paragraphs[4].text = "合成篡改"
            save(edited, path)

        with patch.object(WordDocument, "save", corrupt), self.assertRaisesRegex(DocumentError, "回读"):
            self.write(write_docx, "docx")

    def test_markdown_body_is_inert_and_reversible(self):
        """验证Markdown正文安全编码并准确还原。"""
        text = "    # 标题\n![图片](https://example.invalid/a) <script>x</script>\n`命令` &amp; **加粗**\t\r\n"
        path = self.write(write_markdown, "md", sample(text))
        content = path.read_text(encoding="utf-8")
        marker, escaped = content.split("### ", 1)[1].split("\n\n", 1)
        self.assertIn("00:00:00.123 — 01:00:00.789", marker)
        self.assertEqual(self.decode_markdown(escaped[:-1]), text)
        self.assertNotIn("![", escaped)
        self.assertNotIn("<script>", escaped)
        self.assertNotIn("https://", escaped)
        self.assertNotIn("    #", escaped)
        self.assertNotRegex(escaped, r"(?<!\\)`")

    def test_markdown_does_not_reinterpret_original_entities(self):
        """验证Markdown保留原始HTML实体文本。"""
        text = "&lt;原文&gt; &#65; \\反斜线"
        path = self.write(write_markdown, "md", sample(text))
        escaped = path.read_text(encoding="utf-8").split("### ", 1)[1].split("\n\n", 1)[1][:-1]
        self.assertEqual(self.decode_markdown(escaped), text)

    @staticmethod
    def decode_markdown(text):
        """还原Markdown转义以核对原始正文。"""
        return html.unescape(re.sub(r"\\([" + re.escape(string.punctuation) + r"])", r"\1", text))

    def test_markdown_preserves_normal_spaces_for_readability(self):
        """验证Markdown保留普通空格以便阅读。"""
        text = "    Leading spaces\nPlain English is readable with spaces."
        path = self.write(write_markdown, "md", sample(text))
        content = path.read_text(encoding="utf-8")
        self.assertIn("Plain English is readable with spaces", content)
        self.assertIn("&#32;&#32;&#32;&#32;Leading spaces", content)

    def test_all_formats_support_atomic_publication_temporary_suffix(self):
        """验证三格式支持发布前的临时文件后缀。"""
        for writer in (write_xlsx, write_docx, write_markdown):
            with self.subTest(writer=writer.__name__):
                self.write(writer, "partial")

    def test_all_formats_fail_safely_for_missing_destination(self):
        """验证三格式在目录缺失时安全报错。"""
        for writer in (write_xlsx, write_docx, write_markdown):
            with self.subTest(writer=writer.__name__), self.assertRaises(DocumentError):
                writer(sample(), self.runtime.path("missing/output.bin"),
                       source_name="test.wav", job_id="a" * 32)
