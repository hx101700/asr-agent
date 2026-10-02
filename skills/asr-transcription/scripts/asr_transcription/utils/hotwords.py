"""读取热词Excel原始行并生成词表模板。"""

import io
import zipfile
from pathlib import Path
from xml.etree.ElementTree import ParseError

from defusedxml.common import DefusedXmlException
from openpyxl import Workbook, load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from ..models import HotwordRow


# 本机工作簿解析资源上限，不是阿里云接口限制。
MAX_XLSX_BYTES = 5_000_000
MAX_XLSX_UNCOMPRESSED_BYTES = 20 * 1024 * 1024
MAX_XLSX_ENTRIES = 200
MAX_WORKSHEET_ROWS = 10_001


class HotwordFileError(ValueError):
    """表示热词Excel文件结构错误及行级提示。"""

    def __init__(self, message: str, details: list[dict] | None = None):
        """保存可展示的文件错误及可选行级修改提示。"""
        super().__init__(message)
        self.details = details or []


def _check_xlsx_archive(path: Path) -> None:
    """检查Excel压缩包大小、展开规模和加密标记。"""
    if path.stat().st_size > MAX_XLSX_BYTES:
        raise HotwordFileError("热词文件超过 5 MB 文件上限。")
    with zipfile.ZipFile(path) as archive:
        entries = archive.infolist()
        if (len(entries) > MAX_XLSX_ENTRIES
                or sum(item.file_size for item in entries) > MAX_XLSX_UNCOMPRESSED_BYTES):
            raise HotwordFileError("热词Excel结构过大：解压内容最多20 MiB、200个内部文件。")
        if any(item.flag_bits & 1 for item in entries):
            raise HotwordFileError("不支持加密的热词Excel，请保存为普通.xlsx文件。")


def read_hotwords(path: Path) -> tuple[list[HotwordRow], list[str]]:
    """读取固定两列原始值，返回词条行及工作簿提示。"""
    workbook = None
    try:
        _check_xlsx_archive(path)
        # 有界的小型工作簿直接读取，不依赖可伪造的dimension标签。
        # 显式拥有句柄，XML解析失败时也立即关闭，便于Windows清理副本。
        with path.open("rb") as stream:
            workbook = load_workbook(stream, read_only=False, data_only=False, keep_links=False)
        if "热词" in workbook.sheetnames:
            sheet = workbook["热词"]
        elif len(workbook.sheetnames) == 1:
            sheet = workbook.active
        else:
            raise HotwordFileError("多个工作表时请将待使用的工作表命名为“热词”。")
        if sheet.max_row > MAX_WORKSHEET_ROWS or sheet.max_column > 2:
            raise HotwordFileError("热词工作表仅支持两列、最多10001行（含表头和空行）。")
        header = [sheet.cell(1, number).value for number in (1, 2)]
        if header[0] not in ("text", "热词") or header[1] not in ("weight", "权重"):
            raise HotwordFileError("首行必须依次为text、weight，或中文列名热词、权重。",
                                   [{"row": 1, "field": "header", "message": "请按模板修改表头。"}])
        warnings = []
        if len(workbook.sheetnames) > 1:
            warnings.append("仅读取名为“热词”的工作表，其他工作表不参与此次转写。")
        if header != ["text", "weight"]:
            warnings.append("已按中文别名读取表头：热词对应text，权重对应weight。")
        rows = [HotwordRow(number, cells[0].value, cells[1].value,
                           any(cell.data_type == "f" for cell in cells))
                for number, cells in enumerate(sheet.iter_rows(min_row=2, max_col=2), start=2)]
        return rows, warnings
    except (OSError, zipfile.BadZipFile, InvalidFileException, ParseError,
            DefusedXmlException, KeyError, ValueError) as exc:
        if isinstance(exc, HotwordFileError):
            raise
        raise HotwordFileError("无法读取热词Excel，请检查文件是否损坏或仍在写入。") from exc
    finally:
        if workbook is not None:
            workbook.close()


def hotwords_template() -> bytes:
    """生成带示例词条的两列Excel模板，供网页直接下载。"""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "热词"
    sheet.append(["text", "weight"])
    sheet.append(["示例术语", 4])
    sheet.column_dimensions["A"].width = 32
    sheet.column_dimensions["B"].width = 12
    sheet.freeze_panes = "A2"
    buffer = io.BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()
