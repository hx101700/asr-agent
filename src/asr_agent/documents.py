"""本地转写成品写入器；每种格式写入后独立核验，不接触云端。"""

import html
import math
import re
import string
import unicodedata
from datetime import timedelta
from pathlib import Path
from zipfile import BadZipFile

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from . import MODEL
from .results import Sentence, Transcript, timestamp

TITLE = "录音转写"
SHEET = "转写明细"
HEADERS = ("序号", "音轨", "开始时间", "结束时间", "说话人", "转写内容")
FIRST_ROW = 6
TIME_FORMAT = "[hh]:mm:ss.000"
INVALID_XML = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]")


class DocumentError(ValueError):
    """可向用户展示的格式写入错误，不包含转写正文。"""


def _speaker(sentence: Sentence) -> str:
    return "未标注" if sentence.speaker_id is None else f"说话人{sentence.speaker_id}"


def _track(sentence: Sentence) -> str:
    if sentence.channel_id is None:
        return f"未标注（第{sentence.track_index}组）"
    return f"音轨{sentence.channel_id}"


def _label(sentence: Sentence) -> str:
    return (f"{sentence.index} · {timestamp(sentence.begin_ms)} — {timestamp(sentence.end_ms)}"
            f" · {_speaker(sentence)} · {_track(sentence)}")


def _xml_text(text: str, format_name: str) -> None:
    # XML不能表示这些控制字符；拒绝该格式，不能悄悄删除识别正文中的字符。
    if INVALID_XML.search(text):
        raise DocumentError(f"{format_name}无法保存结果中的控制字符；原始JSON及其他格式不受影响。")


def _cell_value(cell, value) -> None:
    if isinstance(value, str):
        _xml_text(value, "Excel")
        if len(value) > 32767:
            raise DocumentError("Excel单元格最多保存32767个字符，本次未截断内容。请查看JSON或其他成品。")
        cell.value = value
        # 转写和文件名都是不可信文本；即使以=开头也不能成为Excel公式。
        cell.data_type = "s"
    else:
        cell.value = value


def _row_height(text: str, width: float) -> float:
    lines = 0
    for line in text.replace("\r\n", "\n").replace("\r", "\n").split("\n"):
        units = sum(4 if char == "\t" else 2.2 if unicodedata.east_asian_width(char) in "WF"
                    else 1.1 for char in line)
        lines += max(1, math.ceil(units / (width - 4)))
    height = max(30, 20 * lines + 10)
    if height > 409:
        raise DocumentError("单段内容超过Excel行高展示范围，本次未裁切内容。请查看JSON、Word或Markdown。")
    return height


def _xlsx_values(sentence: Sentence) -> tuple:
    return (sentence.index, _track(sentence), timedelta(milliseconds=sentence.begin_ms),
            timedelta(milliseconds=sentence.end_ms), _speaker(sentence), sentence.text)


def write_xlsx(transcript: Transcript, path: Path, *, source_name: str, job_id: str) -> None:
    """时间保存为可计算的Excel时长；所有文字按字符串保存并回读核验。"""
    if len(transcript.sentences) + FIRST_ROW - 1 > 1_048_576:
        raise DocumentError("转写段落超过Excel工作表行数上限，本次未截断内容。")
    workbook = Workbook()
    try:
        sheet = workbook.active
        sheet.title = SHEET
        sheet.sheet_view.showGridLines = False
        sheet.sheet_view.zoomScale = 85
        widths = (8, 22, 18, 18, 18, 88)
        for column, width in zip("ABCDEF", widths):
            sheet.column_dimensions[column].width = width
        metadata = (TITLE, f"音频：{source_name}", f"模型：{MODEL}    任务：{job_id}")
        for row, value in enumerate(metadata, 1):
            sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
            cell = sheet.cell(row, 1)
            _cell_value(cell, value)
            cell.font = Font(name="Microsoft YaHei", size=16 if row == 1 else 10,
                             bold=row == 1, color="000000" if row == 1 else "475467")
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            sheet.row_dimensions[row].height = 38 if row == 1 else _row_height(value, sum(widths))
        sheet.row_dimensions[4].height = 10
        for column, value in enumerate(HEADERS, 1):
            cell = sheet.cell(FIRST_ROW - 1, column)
            _cell_value(cell, value)
            cell.font = Font(name="Microsoft YaHei", size=11, bold=True, color="344054")
            cell.fill = PatternFill("solid", fgColor="FFF2E8")
            cell.alignment = Alignment(horizontal="center", vertical="center")
        sheet.row_dimensions[FIRST_ROW - 1].height = 30
        body_font = Font(name="Microsoft YaHei", size=11, color="101828")
        bottom = Border(bottom=Side(style="hair", color="E4E7EC"))
        for row, sentence in enumerate(transcript.sentences, FIRST_ROW):
            values = _xlsx_values(sentence)
            for column, value in enumerate(values, 1):
                cell = sheet.cell(row, column)
                _cell_value(cell, value)
                cell.font = body_font
                cell.border = bottom
                cell.alignment = Alignment(horizontal="right" if column in (1, 3, 4) else "left",
                                           vertical="top", wrap_text=True, indent=1)
                if row % 2:
                    cell.fill = PatternFill("solid", fgColor="F8FAFC")
                if column in (3, 4):
                    cell.number_format = TIME_FORMAT
            sheet.row_dimensions[row].height = max(
                _row_height(str(value), width)
                for value, width in zip((sentence.index, _track(sentence), timestamp(sentence.begin_ms),
                                         timestamp(sentence.end_ms), _speaker(sentence), sentence.text), widths)
            )
        sheet.freeze_panes = "B6"
        sheet.auto_filter.ref = f"A5:F{sheet.max_row}"
        sheet.print_title_rows = "1:5"
        sheet.print_area = f"A1:F{sheet.max_row}"
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        workbook.save(path)
    except (OSError, ValueError, BadZipFile) as exc:
        if isinstance(exc, DocumentError):
            raise
        raise DocumentError("Excel写入失败，请检查保存位置和文件占用；未自动重试。") from exc
    finally:
        workbook.close()
    try:
        # 以文件对象读取，兼容应用层用于原子发布的临时扩展名。
        with path.open("rb") as stream:
            verified = load_workbook(stream, read_only=True, data_only=False)
            try:
                sheet = verified[SHEET]
                if (sheet.max_row != len(transcript.sentences) + FIRST_ROW - 1
                        or sheet.max_column != 6 or sheet["A2"].value != metadata[1]):
                    raise DocumentError("Excel回读的段落数或元信息不一致，未报告完成。")
                rows = sheet.iter_rows(min_row=FIRST_ROW, max_col=6)
                for cells, sentence in zip(rows, transcript.sentences, strict=True):
                    values = tuple(cell.value if cell.value is not None else "" for cell in cells)
                    if (values != _xlsx_values(sentence)
                            or any(cell.data_type == "f" for cell in cells)):
                        raise DocumentError("Excel回读的正文、时间或标签不一致，未报告完成。")
            finally:
                verified.close()
    except (OSError, ValueError, KeyError, BadZipFile) as exc:
        if isinstance(exc, DocumentError):
            raise
        raise DocumentError("Excel无法重新打开核验，未报告完成。") from exc


def _word_text(paragraph, text: str) -> None:
    _xml_text(text, "Word")
    # python-docx把CR转换成LF。独立写入CR文本节点，以便回读时仍保留原始字符。
    for index, part in enumerate(text.split("\r")):
        if index:
            element = OxmlElement("w:t")
            element.text = "\r"
            paragraph.add_run()._r.append(element)
        paragraph.add_run(part)


def _word_style(style, *, size: float, bold: bool = False) -> None:
    style.font.name = "Microsoft YaHei"
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")


def write_docx(transcript: Transcript, path: Path, *, source_name: str, job_id: str) -> None:
    """用段落承载正文，允许跨页；时间标签与下一段绑定，避免孤立标签。"""
    expected = [TITLE, f"音频：{source_name}", f"模型：{MODEL}", f"任务：{job_id}"]
    try:
        document = Document()
        section = document.sections[0]
        section.page_width, section.page_height = Inches(8.5), Inches(11)
        section.top_margin = section.bottom_margin = Inches(0.8)
        section.left_margin = section.right_margin = Inches(0.85)
        _word_style(document.styles["Normal"], size=11)
        body_format = document.styles["Normal"].paragraph_format
        body_format.line_spacing = 1.35
        body_format.space_after = Pt(9)
        body_format.keep_together = False
        body_format.widow_control = True
        _word_style(document.styles["Title"], size=24, bold=True)
        document.styles["Title"].paragraph_format.space_after = Pt(14)
        metadata_style = document.styles.add_style("Transcript Metadata", WD_STYLE_TYPE.PARAGRAPH)
        metadata_style.base_style = document.styles["Normal"]
        _word_style(metadata_style, size=9)
        metadata_style.paragraph_format.space_after = Pt(4)
        marker_style = document.styles.add_style("Transcript Marker", WD_STYLE_TYPE.PARAGRAPH)
        marker_style.base_style = document.styles["Normal"]
        _word_style(marker_style, size=9, bold=True)
        marker_style.paragraph_format.keep_with_next = True
        marker_style.paragraph_format.space_before = Pt(8)
        marker_style.paragraph_format.space_after = Pt(4)
        for index, text in enumerate(expected):
            paragraph = document.add_paragraph(style="Title" if index == 0 else metadata_style)
            _word_text(paragraph, text)
        for sentence in transcript.sentences:
            label = _label(sentence)
            _word_text(document.add_paragraph(style=marker_style), label)
            _word_text(document.add_paragraph(), sentence.text)
            expected.extend((label, sentence.text))
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer.add_run("第 ")
        page = OxmlElement("w:fldSimple")
        page.set(qn("w:instr"), "PAGE")
        footer._p.append(page)
        footer.add_run(" 页")
        document.core_properties.title = TITLE
        document.core_properties.author = "asr-agent"
        document.save(path)
        with path.open("rb") as stream:
            restored = Document(stream)
            if [paragraph.text for paragraph in restored.paragraphs] != expected:
                raise DocumentError("Word回读的正文、时间或标签不一致，未报告完成。")
    except (OSError, ValueError, KeyError, BadZipFile) as exc:
        if isinstance(exc, DocumentError):
            raise
        raise DocumentError("Word写入或回读失败，请检查保存位置和文件占用；未自动重试。") from exc


def _markdown_text(text: str) -> str:
    # CommonMark反斜杠转义阻止链接/图像等语法；HTML字符另行编码。
    # 只编码行首空白来避免代码块，英文正文保留普通空格，便于直接阅读源文件。
    result = []
    line_start = True
    for char in text:
        if char == "\n":
            result.append(char)
            line_start = True
            continue
        if (line_start and char in " \t") or char == "\r":
            result.append(f"&#{ord(char)};")
            continue
        line_start = False
        if char in "<>&":
            result.append(html.escape(char))
        elif char in string.punctuation:
            result.append("\\" + char)
        else:
            result.append(char)
    return "".join(result)


def _markdown_original(text: str) -> str:
    unescaped = re.sub(r"\\([" + re.escape(string.punctuation) + r"])", r"\1", text)
    return html.unescape(unescaped)


def write_markdown(transcript: Transcript, path: Path, *, source_name: str, job_id: str) -> None:
    parts = [f"# {TITLE}", f"音频：{_markdown_text(source_name)}", f"模型：{_markdown_text(MODEL)}",
             f"任务：{_markdown_text(job_id)}"]
    for sentence in transcript.sentences:
        escaped = _markdown_text(sentence.text)
        if _markdown_original(escaped) != sentence.text:
            raise DocumentError("Markdown无法无损表示结果中的字符，未改写原文。")
        parts.extend((f"### {_label(sentence)}", escaped))
    content = "\n\n".join(parts) + "\n"
    try:
        with path.open("w", encoding="utf-8", newline="") as output:
            output.write(content)
        with path.open("r", encoding="utf-8", newline="") as saved:
            if saved.read() != content:
                raise DocumentError("Markdown回读的正文或时间不一致，未报告完成。")
    except (OSError, UnicodeError) as exc:
        raise DocumentError("Markdown写入或回读失败，请检查保存位置和文件占用；未自动重试。") from exc
