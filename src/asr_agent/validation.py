"""上传前的本地输入检查；不联网、不转换媒体、不修改用户文件。"""

import hashlib
import math
import stat
import zipfile
from dataclasses import asdict
from pathlib import Path
from xml.etree.ElementTree import ParseError

from defusedxml.common import DefusedXmlException
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from .environment import Project
from .media import MediaError, probe_audio


# 官方模型/临时OSS限制见doc/REFERENCES.md A02、A04、A06。
AUDIO_SUFFIXES = frozenset(f".{name}" for name in (
    "aac", "amr", "avi", "flac", "flv", "m4a", "mkv", "mov", "mp3",
    "mp4", "mpeg", "ogg", "opus", "wav", "webm", "wma", "wmv",
))
SUPPORTED_CONTAINERS = frozenset({
    "aac", "amr", "avi", "flac", "flv", "mov", "mp3", "mpeg", "ogg",
    "wav", "matroska", "webm", "asf",
})
MAX_UPLOAD_BYTES = 1_000_000_000
MAX_DURATION_SECONDS = 12 * 60 * 60
MAX_HOTWORDS = 2000
# 以下是本地解析资源上限，不是阿里云接口限制。
MAX_XLSX_BYTES = 5 * 1024 * 1024
MAX_XLSX_UNCOMPRESSED_BYTES = 20 * 1024 * 1024
MAX_XLSX_ENTRIES = 200
MAX_WORKSHEET_ROWS = 10_001
PROTECTED_NAMES = frozenset({"node_modules", "__pycache__", "venv", "secrets", "credentials"})


class ValidationError(ValueError):
    """可直接按字段呈现的本地错误；details只包含行号、字段与修改说明。"""

    def __init__(self, message: str, field: str, details: list[dict] | None = None):
        super().__init__(message)
        self.message = message
        self.field = field
        self.details = details or []


def _check_relative_path(root: Path, path: Path, field: str) -> None:
    if not path.is_relative_to(root):
        raise ValidationError("请选择当前项目目录内的文件。", field)
    parts = path.relative_to(root).parts
    if any(part.startswith(".") or part.casefold() in PROTECTED_NAMES for part in parts):
        raise ValidationError("不能读取隐藏文件、凭据目录或程序内部目录。", field)


def _resolve_local_path(project: Project, value: str | Path, field: str) -> Path:
    """同时检查用户路径与链接目标，避免通过junction读取项目外或隐藏目录。"""
    if not isinstance(value, (str, Path)) or not str(value).strip():
        raise ValidationError("请指定一个本地路径。", field)
    try:
        root = project.root.resolve(strict=True)
        entered = Path(value)
        candidate = entered if entered.is_absolute() else root / entered
        _check_relative_path(root, candidate, field)
        path = candidate.resolve(strict=True)
        _check_relative_path(root, path, field)
        for part in (path, *path.parents):
            if part == root:
                break
            attributes = getattr(part.stat(), "st_file_attributes", 0)
            if attributes & (stat.FILE_ATTRIBUTE_HIDDEN | stat.FILE_ATTRIBUTE_SYSTEM):
                raise ValidationError("不能读取隐藏文件或系统文件。", field)
        return path
    except (OSError, RuntimeError, ValueError) as exc:
        if isinstance(exc, ValidationError):
            raise
        raise ValidationError("无法读取指定文件，请检查路径和访问权限。", field) from exc


def resolve_input(project: Project, value: str | Path, allowed_suffixes, field: str) -> Path:
    path = _resolve_local_path(project, value, field)
    if not path.is_file():
        raise ValidationError("请选择普通文件，不能选择目录。", field)
    if path.suffix.lower() not in allowed_suffixes:
        raise ValidationError("文件扩展名不在允许的格式范围内。", field)
    return path


def file_fingerprint(path: Path) -> dict:
    """分块计算指纹；读取期间发生改动则要求重新校验，避免确认旧内容。"""
    try:
        before = path.stat()
        digest = hashlib.sha256()
        with path.open("rb") as source:
            for chunk in iter(lambda: source.read(1024 * 1024), b""):
                digest.update(chunk)
        after = path.stat()
        if (before.st_size, before.st_mtime_ns, before.st_ino) != (
            after.st_size, after.st_mtime_ns, after.st_ino,
        ):
            raise ValidationError("文件在校验期间发生变化，请重新校验。", "file")
        return {"size_bytes": after.st_size, "mtime_ns": after.st_mtime_ns,
                "sha256": digest.hexdigest()}
    except OSError as exc:
        raise ValidationError("无法读取文件，请检查访问权限后重新校验。", "file") from exc


def _check_unchanged(path: Path, fingerprint: dict, field: str) -> None:
    current = path.stat()
    if (current.st_size, current.st_mtime_ns) != (
        fingerprint["size_bytes"], fingerprint["mtime_ns"],
    ):
        raise ValidationError("文件在校验期间发生变化，请重新校验。", field)


def validate_audio(project: Project, path: str | Path, diarization: bool) -> dict:
    """只探测原文件；需要合并时留到确认后，并再次校验实际上传副本。"""
    field = "audio_path"
    source = resolve_input(project, path, AUDIO_SUFFIXES, field)
    if not isinstance(diarization, bool):
        raise ValidationError("说话人选项必须为开启或关闭。", "diarization")
    try:
        before = source.stat()
        if before.st_size == 0:
            raise ValidationError("音频文件为空。", field)
        info = probe_audio(source)
        if not set(info.format_name.split(",")) & SUPPORTED_CONTAINERS:
            raise ValidationError("实际媒体格式不在固定模型支持范围内。", field)
        duration = info.duration_seconds
        if duration is None or not math.isfinite(duration) or duration <= 0:
            raise ValidationError("无法确定有效音频时长，不能完成上传前校验。", field)
        if duration > MAX_DURATION_SECONDS:
            raise ValidationError("音频时长超过模型允许的12小时。", field)
        requires_mono = diarization and info.channels > 1
        if not requires_mono and info.size_bytes > MAX_UPLOAD_BYTES:
            raise ValidationError("待上传音频超过临时OSS的1 GB上限。", field)
        warnings = []
        if requires_mono:
            warnings.append(
                f"此音频包含 {info.channels} 个声道。为区分发言人，转写前将生成单声道 FLAC 副本，"
                "保留原文件。副本通过大小和时长检查后才会上传。"
            )
        if diarization and duration > 2 * 60 * 60:
            warnings.append("音频超过 2 小时。启用发言人区分可能导致识别失败或超时，建议使用 2 小时以内的音频。")
        if info.audio_tracks > 1:
            warnings.append(f"此文件包含 {info.audio_tracks} 个音轨，仅转写第一个音轨（索引0），其余音轨不会转写。")
        fingerprint = file_fingerprint(source)
        after = source.stat()
        if (fingerprint["size_bytes"] != info.size_bytes
                or (before.st_size, before.st_mtime_ns, before.st_ino)
                != (after.st_size, after.st_mtime_ns, after.st_ino)):
            raise ValidationError("文件在校验期间发生变化，请重新校验。", field)
        return {"path": str(source), "metadata": asdict(info), "fingerprint": fingerprint,
                "requires_mono": requires_mono, "warnings": warnings}
    except (MediaError, OSError) as exc:
        raise ValidationError("无法读取音频，请检查文件是否损坏及格式是否支持。", field) from exc


def validate_context(text: str) -> str:
    if not isinstance(text, str) or not text.strip():
        raise ValidationError("请输入参考文本，或关闭上下文增强。", "context")
    if len(text) > 400:
        raise ValidationError(f"参考文本共 {len(text)} 个字符，最多支持 400 个，请精简后重新检查。", "context")
    if "\x00" in text or any(0xD800 <= ord(char) <= 0xDFFF for char in text):
        raise ValidationError("参考文本中含有无法识别的特殊字符，请删除后重新检查。", "context")
    return text


def _check_xlsx_archive(path: Path) -> None:
    if path.stat().st_size > MAX_XLSX_BYTES:
        raise ValidationError("热词Excel超过本地5 MiB文件上限。", "hotwords_path")
    with zipfile.ZipFile(path) as archive:
        entries = archive.infolist()
        if (len(entries) > MAX_XLSX_ENTRIES
                or sum(item.file_size for item in entries) > MAX_XLSX_UNCOMPRESSED_BYTES):
            raise ValidationError("热词Excel结构过大：解压内容最多20 MiB、200个内部文件。", "hotwords_path")
        if any(item.flag_bits & 1 for item in entries):
            raise ValidationError("不支持加密的热词Excel，请保存为普通.xlsx文件。", "hotwords_path")


def load_hotwords(project: Project, path: str | Path, allow_super: bool = False) -> dict:
    """读取固定两列Excel，错误逐行返回；不执行公式、不默默修剪或丢弃词条。"""
    field = "hotwords_path"
    source = resolve_input(project, path, {".xlsx"}, field)
    workbook = None
    try:
        _check_xlsx_archive(source)
        fingerprint = file_fingerprint(source)
        # 有界的小型工作簿直接读取，避免依赖可伪造的只读模式dimension标签。
        workbook = load_workbook(source, read_only=False, data_only=False, keep_links=False)
        if "热词" in workbook.sheetnames:
            sheet = workbook["热词"]
        elif len(workbook.sheetnames) == 1:
            sheet = workbook.active
        else:
            raise ValidationError("多个工作表时请将待使用的工作表命名为“热词”。", field)
        if sheet.max_row > MAX_WORKSHEET_ROWS or sheet.max_column > 2:
            raise ValidationError("热词工作表仅支持两列、最多10001行（含表头和空行）。", field)
        header = [sheet.cell(1, number).value for number in (1, 2)]
        if header[0] not in ("text", "热词") or header[1] not in ("weight", "权重"):
            raise ValidationError("首行必须依次为text、weight，或中文列名热词、权重。", field,
                                  [{"row": 1, "field": "header", "message": "请按模板修改表头。"}])
        warnings = []
        if len(workbook.sheetnames) > 1:
            warnings.append("仅读取名为“热词”的工作表，其他工作表不参与此次转写。")
        if header != ["text", "weight"]:
            warnings.append("已按中文别名读取表头：热词对应text，权重对应weight。")
        vocabulary = {}
        first_rows = {}
        details = []
        ignored_blank_rows = 0
        super_count = 0
        for row_number, cells in enumerate(sheet.iter_rows(min_row=2, max_col=2), start=2):
            values = [cell.value for cell in cells]
            if all(value is None for value in values):
                ignored_blank_rows += 1
                continue
            if any(cell.data_type == "f" for cell in cells):
                details.append({"row": row_number, "field": "row", "message": "不接受公式，请填写固定文本和数值。"})
                continue
            text, weight = values
            row_errors = []
            if not isinstance(text, str) or not text.strip():
                row_errors.append(("text", "热词必须为非空文本。"))
            elif text != text.strip() or any(ord(char) < 32 or ord(char) == 127 for char in text):
                row_errors.append(("text", "请移除热词首尾空白、换行或控制字符；程序不会自动修改。"))
            elif not text.isascii() and len(text) > 15:
                row_errors.append(("text", "含非ASCII字符时，热词总长度最多15个字符。"))
            elif text.isascii() and len([part for part in text.split(" ") if part]) > 7:
                row_errors.append(("text", "纯ASCII热词按空格切分后最多7段。"))
            allowed_weights = (1, 2, 3, 4, 5, 50) if allow_super else (1, 2, 3, 4, 5)
            if (isinstance(weight, bool) or not isinstance(weight, (int, float))
                    or weight not in allowed_weights):
                message = "权重必须为1至5的整数" + ("或50。" if allow_super else "。")
                row_errors.append(("weight", message))
            if row_errors:
                details.extend({"row": row_number, "field": name, "message": message}
                               for name, message in row_errors)
                continue
            weight = int(weight)
            if text in vocabulary:
                if vocabulary[text] != weight:
                    details.append({"row": row_number, "field": "weight",
                                    "message": f"与第{first_rows[text]}行热词重复但权重不同，请统一。"})
                else:
                    warnings.append(f"第{row_number}行与第{first_rows[text]}行完全相同，已合并为一个词条。")
                continue
            vocabulary[text] = weight
            first_rows[text] = row_number
            if len(vocabulary) == MAX_HOTWORDS + 1:
                details.append({"row": row_number, "field": "text", "message": "热词总数超过2000个，请减少。"})
            if weight == 50:
                super_count += 1
                if super_count == 51:
                    details.append({"row": row_number, "field": "weight", "message": "超级热词（权重50）最多50个。"})
        if details:
            raise ValidationError("热词文件中有不符合要求的内容，请按以下行号修改后重新添加文件。", field, details)
        if not vocabulary:
            raise ValidationError("热词Excel未包含有效词条。", field)
        if ignored_blank_rows:
            warnings.append(f"已忽略{ignored_blank_rows}个完全空白行。")
        _check_unchanged(source, fingerprint, field)
        return {"path": str(source), "sheet": sheet.title, "vocabulary": vocabulary,
                "count": len(vocabulary), "super_count": super_count,
                "ignored_blank_rows": ignored_blank_rows, "fingerprint": fingerprint,
                "warnings": warnings}
    except (OSError, zipfile.BadZipFile, InvalidFileException, ParseError,
            DefusedXmlException, KeyError, ValueError) as exc:
        if isinstance(exc, ValidationError):
            raise
        raise ValidationError("无法读取热词Excel，请检查文件是否损坏或仍在写入。", field) from exc
    finally:
        if workbook is not None:
            workbook.close()
